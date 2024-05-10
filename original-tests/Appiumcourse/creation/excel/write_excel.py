import openpyxl

# Create a new workbook and set the active sheet
workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = "Sheet1"

# Write header in first row
sheet.cell(row=1, column=1).value = "Name"
sheet.cell(row=1, column=2).value = "Email"
sheet.cell(row=1, column=3).value = "Age"

# Fill data in rows 4-7 and columns 1-3
for row in range(4, 8):
    for col in range(1, 4):
        # Using the loop variables to write to different cells
        if col == 1:
            sheet.cell(row=row, column=col).value = f"Person {row-3}"
        elif col == 2:
            sheet.cell(row=row, column=col).value = f"person{row-3}@example.com"
        elif col == 3:
            sheet.cell(row=row, column=col).value = 20 + (row-3)

# Save the workbook after all data is written
workbook.save("test.xlsx")
print("Excel file 'test.xlsx' created successfully with data.")

